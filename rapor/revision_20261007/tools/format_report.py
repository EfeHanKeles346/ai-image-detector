"""Format a student-saved DOCX without rewriting its text; optional rendered TOC map."""
from pathlib import Path
import argparse, json, hashlib, re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

parser=argparse.ArgumentParser()
parser.add_argument('input');parser.add_argument('output');parser.add_argument('--pages')
a=parser.parse_args();d=Document(a.input)
# Capture every non-empty, non-TOC text paragraph, including captions and tables.
def contents(doc):
 return [''.join(p.xpath('.//w:t/text()')) for p in doc.element.body.xpath('.//w:p[not(ancestor::w:sdt)]') if ''.join(p.xpath('.//w:t/text()')).strip()]
before=contents(d)
for sec in d.sections:
 sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
for style in d.styles:
 if style.type in (1,2):
  style.font.name='Times New Roman';style.font.size=Pt(12);style.font.color.rgb=RGBColor(0,0,0)
for name in ['Normal','Title','Subtitle','Caption','Heading 1','Heading 2','Heading 3']:
 f=d.styles[name].paragraph_format;f.left_indent=f.right_indent=f.first_line_indent=Pt(0)
 f.line_spacing=2;f.space_before=f.space_after=Pt(0)
 f.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
for name in ['Title','Heading 1','Heading 2','Heading 3']:
 d.styles[name].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
 d.styles[name].font.bold=True
 d.styles[name].paragraph_format.keep_with_next=True
refs=False;removed=0;in_body=False
for p in list(d.paragraphs):
 t=p.text.strip();f=p.paragraph_format
 if t=='1. Introduction':in_body=True
 if in_body:
  for br in p._p.xpath('.//w:br[@w:type="page"]'):br.getparent().remove(br)
 if not t:
  if p._p.xpath('.//w:drawing'):
   p.alignment=WD_ALIGN_PARAGRAPH.CENTER;f.line_spacing=1;f.keep_with_next=True
   f.left_indent=f.right_indent=f.first_line_indent=Pt(0)
  elif not p._p.xpath('.//w:br|.//w:fldChar|.//w:instrText|./w:pPr/w:sectPr'):
   p._p.getparent().remove(p._p);removed+=1
  continue
 f.left_indent=f.right_indent=f.first_line_indent=Pt(0)
 f.line_spacing=2;f.space_before=f.space_after=Pt(0);f.widow_control=True
 p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
 if p.style.name.startswith('Heading'):
  refs=t=='8. References' if p.style.name=='Heading 1' else refs
  p.alignment=WD_ALIGN_PARAGRAPH.CENTER if t=='8. References' else WD_ALIGN_PARAGRAPH.LEFT
  f.space_before=Pt(12);f.space_after=Pt(6);f.keep_with_next=True;f.keep_together=True
  f.page_break_before=t in ['1. Introduction','4.6 Results','8. References','9. Appendices']
 elif p.style.name=='Title':
  p.alignment=WD_ALIGN_PARAGRAPH.LEFT;f.keep_with_next=True
  if t.startswith('PixelProof'):f.space_before=Pt(54);f.space_after=Pt(36)
 elif p.style.name=='Caption':
  p.alignment=WD_ALIGN_PARAGRAPH.CENTER;f.keep_together=True
  f.keep_with_next=t.startswith('Table ')
 elif refs:
  f.left_indent=Pt(15);f.first_line_indent=Pt(-15);f.line_spacing=1;f.space_after=Pt(12);f.keep_together=True
for table in d.tables:
 table.alignment=WD_TABLE_ALIGNMENT.CENTER
 for row in table.rows:
  for cell in row.cells:
   for p in cell.paragraphs:
    p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    f=p.paragraph_format;f.left_indent=f.right_indent=f.first_line_indent=Pt(0)
    f.line_spacing=2;f.space_before=f.space_after=Pt(0)
for shape in d.inline_shapes:
 for k in ['distT','distB','distL','distR']:shape._inline.set(k,'0')
 ext=shape._inline.find(qn('wp:effectExtent'))
 if ext is None:ext=OxmlElement('wp:effectExtent');shape._inline.insert(1,ext)
 for k in ['l','t','r','b']:ext.set(k,'0')
# Preserve the TOC field, updating its cached page values only after rendering.
page_map=json.loads(Path(a.pages).read_text()) if a.pages else {}
for x in d.element.xpath('.//w:sdt//w:p'):
 p=Paragraph(x,d._body);f=p.paragraph_format
 p.alignment=WD_ALIGN_PARAGRAPH.LEFT;f.left_indent=f.right_indent=f.first_line_indent=Pt(0)
 f.line_spacing=1.1;f.space_before=f.space_after=Pt(0)
 ts=x.xpath('.//w:t')
 if len(ts)>=2 and ts[0].text in page_map:ts[-1].text=str(page_map[ts[0].text])
# Normalize all text, including hyperlinks, fields and table cells, without replacing runs.
parts=[d.element,d.styles.element]
for sec in d.sections:
 parts.extend([sec.header._element,sec.footer._element])
for root in parts:
 for jc in root.xpath('.//w:jc'):
  if jc.get(qn('w:val'))=='start':jc.set(qn('w:val'),'left')
  if jc.get(qn('w:val'))=='end':jc.set(qn('w:val'),'right')
 for rp in root.xpath('.//w:rPr'):
  for tag in ['color','rFonts','sz','szCs']:
   el=rp.find(qn('w:'+tag))
   if el is None:el=OxmlElement('w:'+tag);rp.append(el)
   el.attrib.clear()
   if tag=='color':el.set(qn('w:val'),'000000')
   elif tag=='rFonts':
    for attr in ['ascii','hAnsi','eastAsia','cs']:el.set(qn('w:'+attr),'Times New Roman')
   else:el.set(qn('w:val'),'24')
  for x in list(rp):
   if x.tag in [qn('w:highlight'),qn('w:shd')]:rp.remove(x)
 # Runs lacking rPr also need explicit styling.
 for run in root.xpath('.//w:r[not(w:rPr)]'):
  rp=OxmlElement('w:rPr');run.insert(0,rp)
  for tag,val in [('color','000000'),('sz','24')]:
   el=OxmlElement('w:'+tag);el.set(qn('w:val'),val);rp.append(el)
  el=OxmlElement('w:rFonts')
  for attr in ['ascii','hAnsi','eastAsia','cs']:el.set(qn('w:'+attr),'Times New Roman')
  rp.append(el)
assert before==contents(d),'Formatting must preserve all non-TOC text'
Path(a.output).parent.mkdir(parents=True,exist_ok=True);d.save(a.output)
print(json.dumps({'paragraphs_preserved':len(before),'removed_empty_paragraphs':removed,'toc_entries_updated':len(page_map),'output':a.output}))
