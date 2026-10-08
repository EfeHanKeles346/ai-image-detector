from pathlib import Path
from docx import Document
from pypdf import PdfReader
import json,sys
r=Path(__file__).resolve().parents[1];n='CS395_FinalReport_Efe_Han_Keleş_4October2026'
docpath=r/'deliverables'/(n+'.docx');d=Document(docpath);pdf=PdfReader(sys.argv[1]);heads=[p.text for p in d.paragraphs if p.style.name.startswith('Heading')];mp={}
for i,page in enumerate(pdf.pages):
 if i<3:continue
 t=' '.join(page.extract_text().split())
 for h in heads:
  if h in t:mp[h]=i+1
assert len(mp)==len(heads)==31
rows=[p for p in d.paragraphs if '\t' in p.text and p.text[0].isdigit()]
assert len(rows)==31
for p,h in zip(rows,heads):
 ts=p._p.xpath('.//w:t');assert len(ts)==2
 ts[0].text=h;ts[1].text=str(mp[h])
d.save(docpath);(r/'sources/heading_pages.json').write_text(json.dumps(mp,ensure_ascii=False,indent=2)+'\n')
print('pages',len(pdf.pages),'heading_pages',mp)
