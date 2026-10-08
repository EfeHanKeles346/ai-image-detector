from pathlib import Path
import json,re,hashlib
from docx import Document
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH as A
from pypdf import PdfReader
import pdfplumber
R=Path(__file__).resolve().parents[1];N='CS395_FinalReport_Efe_Han_Keleş_4October2026';d=Document(R/'deliverables'/(N+'.docx'))
checks=[]
def ck(n,v):
 checks.append({'check':n,'passed':bool(v)})
 if not v:raise AssertionError(n)
expected=json.loads((R/'sources/expected_prose.json').read_text());actual=[];current=None
for p in d.paragraphs:
 if p.style.name.startswith('Heading'):current=p.text
 elif current and current!='8 REFERENCES' and p.style.name=='Normal' and p.text and p.text!='[TO COMPLETE]':actual.append({'section':current,'text':p.text})
ck('all_56_supplied_body_paragraphs_exact_after_authorized_numbering_and_layout',actual==expected)
ck('two_visible_completion_markers',sum(p.text=='[TO COMPLETE]' for p in d.paragraphs)==2)
ck('figures_and_tables_retained',len(d.inline_shapes)==3 and len(d.tables)==5)
ck('margins',all(abs(m.inches-1)<.001 for s in d.sections for m in [s.left_margin,s.right_margin,s.top_margin,s.bottom_margin]))
refs=False;hs=[]
for p in d.paragraphs:
 t=p.text;f=p.paragraph_format
 if p.style.name.startswith('Heading'):
  hs.append(t);refs=t=='8 REFERENCES' if p.style.name=='Heading 1' else refs
  ck('heading_alignment:'+t,p.alignment==(A.CENTER if refs else A.LEFT))
  ck('numbered_heading:'+t,bool(re.match(r'^\d+(?:\.\d+)*',t)))
 elif refs and t:
  ck('hanging_reference:'+t[:25],f.left_indent.pt==15 and f.first_line_indent.pt==-15 and f.line_spacing==1 and f.space_after.pt==12)
 elif t and p.style.name=='Normal':
  ck('body_format:'+t[:28],p.alignment==A.JUSTIFY and f.line_spacing==2 and all(v.pt==0 for v in [f.left_indent,f.right_indent,f.first_line_indent,f.space_before,f.space_after]))
 elif p.style.name=='Caption':
  ck('caption_centered:'+t[:9],p.alignment==A.CENTER)
  ck('caption_placement:'+t[:9],bool(p._p.getprevious().xpath('.//w:drawing')) if t.startswith('Figure') else p._p.getnext().tag==qn('w:tbl'))
old=Document(R.parent/'revision_20261007/deliverables'/(N+'.docx'))
oldrefs=[p.text for p in old.paragraphs if p.text.startswith(('Abdelhamed,','Adobe. (','Chen, C.','Dwork, C.','Finnvera.','Grommelt,','Guo, C.','Keleş, E.','Kim, D.','Ojha, U.','Oquab, M.','Radford, A.','Türk Telekom.','Wang, S.'))]
newrefs=[p.text for p in d.paragraphs[d.paragraphs.index(next(p for p in d.paragraphs if p.text=='8 REFERENCES'))+1:]] if False else []
seen=False
for p in d.paragraphs:
 if p.text=='8 REFERENCES':seen=True
 elif seen and p.text:newrefs.append(p.text)
ck('15_reference_texts_preserved',newrefs==oldrefs and len(newrefs)==15)
ck('all_table_values_preserved',sorted([str([[c.text for c in row.cells] for row in t.rows]) for t in old.tables])==sorted([str([[c.text for c in row.cells] for row in t.rows]) for t in d.tables]))
ck('no_sentence_lists',not d.element.xpath('.//w:pPr/w:numPr'))
ck('word_TOC_field',any('TOC' in (x.text or '') for x in d.element.xpath('.//w:instrText')))
r=PdfReader(R/'deliverables'/(N+'.pdf'));mp=json.loads((R/'sources/heading_pages.json').read_text());actual_pages={}
for i,p in enumerate(r.pages):
 if i<3:continue
 t=' '.join(p.extract_text().split())
 for h in hs:
  if h in t:actual_pages[h]=i+1
ck('all_31_heading_pages_verified',mp==actual_pages and len(mp)==31)
toc=' '.join(r.pages[2].extract_text().split())
for h,n in mp.items():ck('TOC:'+h,bool(re.search(re.escape(h)+r'\.*\s*'+str(n)+r'\b',toc)))
ck('24_pages',len(r.pages)==24)
with pdfplumber.open(R/'deliverables'/(N+'.pdf')) as pdf:
 ck('black_text',all(c.get('non_stroking_color') in (0,(0,0,0)) for p in pdf.pages for c in p.chars))
 ck('Times_New_Roman_12pt',all('TimesNewRoman' in c['fontname'] and abs(c['size']-12)<.1 for p in pdf.pages for c in p.chars))
 ck('within_margins',all(c['x0']>71 and c['x1']<p.width-71 for p in pdf.pages for c in p.chars if c['text'].strip() and 70<c['top']<p.height-70))
 ck('figures_centered',all(abs((im['x0']+im['x1']-p.width)/2)<1 for p in pdf.pages for im in p.images))
ck('results_one_page',all(' '.join(x['text'].split()) in ' '.join(r.pages[18].extract_text().split()) for x in expected if x['section']=='4.6 Results'))
# Evidence of manuscript issues is reported, never silently fixed.
issues=[
'Abstract has a completion marker; Model2 prose is pending under renumbered 4.5.9. Its preserved Table 4 has no body citation yet.',
'Introduction still says Section 9 contains appendices. Removing that clause would change prose, so it remains pending after the authorized relocation.',
'The criteria paragraph still says lowercase "table 3". Albert requires "Table 3"; capitalization was not authorized.',
'Finnvera (2025) remains in the preserved bibliography but is no longer cited. Nokia statement is attached to the company annual-report citation instead.',
'Radford citation in 3.4 lacks an opening parenthesis, and the Guo discussion is duplicated. Neither punctuation nor duplicate words were removed.',
'5.2 omits the required skill the student wishes university had taught; 5.3 describes two difficulties instead of the requested three; 5.1 does not explicitly state career-plan impact.',
'4.5.2 writes AUC 0.480% instead of 0.480. 4.6 says the previous model missed the image, whereas the missed reference-detected image belongs to E92.',
'4.3 states that a new dataset was added in every experiment; project history does not support that statement. 4.5.1 claims resolution was the sole cause and contrasts data quality with camera features.',
'4.6 population wording repeats AI images and can confuse unique parents with versions; the 160 real plus 160 AI parents are evaluated under two conditions.',
'3.1 supplied mentor address uses Hasan.çontuk whereas prior confirmed contact was hasan.contuk; not corrected without authorization.',
'Raw author note "(abstact 3.2)" was kept in the archived input and excluded as a non-report editing note. Markdown links were displayed as their labels; all prose characters otherwise preserved except approved table numbers.',
'Pending prose may change pagination. Current formatted draft is not submission-ready and no AI score or grade is predicted.'
]
(R/'sources/format_audit.json').write_text(json.dumps({'checks':checks,'passed':len(checks),'pages':len(r.pages),'prose_preserved':True,'content_issues':issues,'visual_review':'pending'},indent=2,ensure_ascii=False)+'\n')
print(json.dumps({'passed':len(checks),'pages':len(r.pages),'issues':len(issues)}))
