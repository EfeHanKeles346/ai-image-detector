from pathlib import Path
from docx import Document
from docx.shared import Pt,Inches,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH as A
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
import hashlib,json,shutil
R=Path(__file__).resolve().parents[1];src=Path('/Users/efehankeles/Desktop/SON RAPOR.docx');snapshot=R/'sources/author_before_corrections.docx'
if not snapshot.exists():shutil.copy2(src,snapshot)
d=Document(snapshot); changes=[]
def replace(i,old,new):
 p=d.paragraphs[i];assert old in p.text,(i,old);before=p.text;p.text=before.replace(old,new);changes.append({'paragraph':i,'before':before,'after':p.text})
def setp(i,new):replace(i,d.paragraphs[i].text,new)
protected={i:d.paragraphs[i].text for i in range(133,137)}
setp(15,"Türk Telekomünikasyon A.Ş. is a corporation which provides a wide range of technological services. The internship was completed at this institution. The internship project named PixelProof had some goals to achieve. One of them was making a detector which can differentiate AI generated images from real images and find areas edited by AI. Another goal was to reduce false alerts, such as classifying an image as AI generated even if it is real. There were several methods used in this project. Experiment results were recorded, and new steps were decided according to these results. The project produced a web demo which includes a system that classifies images. In addition, the web demo was designed with a user-friendly interface. False alerts like classifying a real image as AI generated decreased on the development data. However, this does not indicate that results are fully reliable. There were some challenges. One of them was the difficulty of detecting specific areas manipulated by an AI generator due to data problems and the way AI generators edit images. These problems contributed to this part of the project remaining experimental. For future development, the system should be tested on datasets not used during development, and new actions should be decided according to the results.")
replace(49,'is whether the detected photograph is','is to detect whether a photograph is')
replace(49,'classifies this image as AI generated or real photograph','searches for signs of AI generation')
replace(50,'section 8 contains references and section 9 includes appendices.','section 8 contains references.')
replace(53,'and AssisTT, and one of the main supply is nokia (Türk Telekom, 2026a).','and AssisTT (Türk Telekom, 2026a). Nokia supplies fixed and mobile network equipment and related services to the Türk Telekom group (Finnvera, 2025).')
replace(59,'Hasan.çontuk@turktelekom.com.tr','hasan.contuk@turktelekom.com.tr');replace(59,'Oner Çelebi','Önder Çelebi')
replace(62,'96,75%','96.75%')
replace(62,'These experiments result indicate that different types of datasets, changes the model performance due to format, compression levels and size. After this inference, images were transformed into common format named JPEG and the test was conducted again. There was no clear drop in new test (Keleş 2026).','These results showed that performance changed on a different dataset. Later checks found differences in image format and size between real and AI images. The images were converted to a common JPEG format and cropped into squares, and the test was repeated. There was no clear drop in AUC after these changes, so this check did not prove that those differences caused the original performance gap (Keleş, 2026, E1 and E10).')
replace(64,'For instance, when more AI generated images are added and labeled as ai generated, there might be performance increase in AI detection. However, detection of real images performance might drop down.','For instance, when more images are classified as AI generated, AI detection may increase. However, more real images may also be incorrectly classified as AI generated.')
replace(65,'curial','crucial')
replace(70,'images. (Radford et al., 2021).','images (Radford et al., 2021).')
replace(82,d.paragraphs[82].text,d.paragraphs[82].text+' Production deployment and checking content credentials were outside the scope of this student project.')
replace(84,'feature experiments','future experiments');replace(86,'DATASET.md','DATASETS.md')
replace(88,'calibration values','calibration data');replace(89,'traceable.Full','traceable. Full')
replace(95,'fine tunning','fine tuning') if 'fine tunning' in d.paragraphs[95].text else None
replace(95,'The reason behind this was resolution problem. After fixing this problem performance largely regained.','A resolution control improved performance, suggesting that the difference in input preparation contributed to this result.')
replace(96,'fallowed','followed');replace(100,'preform bad','perform poorly');replace(106,'fine tunning','fine tuning')
replace(111,'This table 3','Table 3')
replace(114,'Figure 3 shows that compare results of two models in 160 real images.','Figure 3 includes E43, E86 and E92 results on the same 160 real images. The following comparison focuses on E43 and E92.')
replace(117,'two different families','two previously seen AI families')
replace(117,'12 out of 160 false alerts in images','12 out of 160 false alerts in processed real images');replace(117,'14 out of 160 in false alerts in real images','14 out of 160 false alerts in processed real images')
replace(117,d.paragraphs[117].text,d.paragraphs[117].text+' (Keleş, 2026, E92 and E102).')
replace(120,'In the previous table,','According to previous findings,')
replace(123,'75.73% to 74.15','75.73% to 74.15%')
replace(125,'The first main goal was Model2 to detect the edited areas which generated by an AI generator.','The main goal of Model2 was to detect areas edited by an AI generator.')
replace(125,'prevented to sustain developmental experiments','limited the development experiments')
replace(125,'which is intendent to manipulate','which was intended to be manipulated')
replace(125,'According to table 4 detecting manipulated area improves, but performance of the original area decreased and false alerts on real images increased.','According to Table 4, detecting edits in a new area improved, but performance in the original area decreased and false alerts on real images increased.')
replace(125,'the related candidates rejected.','this candidate was rejected (Keleş, 2026, E105, E109, E117, E122 and E129–E135).')
replace(129,'Number of real and AI generated images are 160 for original version and 160 for processed version, so in total 320 images for each version. Same situation is valid for AI generated images.','There were 160 real and 160 AI generated images, making 320 initial images. Each had an original and a processed version, producing 640 views in total. These versions came from the same initial images and were not independent examples.')
replace(131,'1.245 Pyhton test executed successfully however this checks the code flow in the system not a measurement of image detection. this system did not used for company use and did not measured impact on the company (Keleş 2026).','1,245 Python tests passed, but these check software behavior, not image detection accuracy. This was a student research prototype. It was not deployed for company use because broader independent evaluation was still needed. Its impact on the company was not measured (Kleş, 2026).')
replace(131,'Kleş','Keleş')
replace(138,'result of the experiment is highly dependent for experiment result','experiment results depend on the quality of the data')
replace(138,d.paragraphs[138].text,d.paragraphs[138].text+' To address this, I checked labels, image quality and shared image sources before using datasets in new experiments. Data that failed these checks were kept out of training.')
replace(140,'I used to worked','I worked');replace(140,'brakes','breaks');replace(140,'I worked 4 p.m.','I worked until 4 p.m.')
replace(143,'but it gives statically consistent results.','but it showed measured improvements on the development data.')
replace(145,'For feature students','For future students');replace(145,'several advise','several recommendations');replace(145,'their inserts','their interests')
for i,old,new in [(48,'1.Introduction:','1. Introduction:'),(51,'2.Company information:','2. Company information:'),(79,'4.Intership project:','4. Internship project:'),(132,'5.Internship experience','5. Internship experience'),(141,'6.Conclusion','6. Conclusion'),(144,'7.Recomandations','7. Recommendations')]:replace(i,old,new)
assert all(d.paragraphs[i].text==t for i,t in protected.items())
# Remove the empty paragraph between Model2 prose and Table 4 caption.
p=d.paragraphs[126];assert not p.text and not p._p.xpath('.//w:pict | .//w:drawing');p._p.getparent().remove(p._p)
# Restore the already approved formatting after edits; no text changes in 5.1/5.2.
formatting=(R.parent/'tools/build_report.py').read_text().split('# Typography:')[1].split('# Preserve all words')[0]
formatting=formatting.replace("'.//w:drawing'","'.//w:drawing | .//w:pict'")
exec('# Typography:'+formatting)
for p in d.paragraphs:
 if '\t' in p.text and p.text[0].isdigit():p.alignment=A.LEFT;p.paragraph_format.line_spacing=1.1
 if p.text=='5. Internship experience':p.paragraph_format.page_break_before=True
for e in d.styles.element.xpath('.//w:numPr'):e.getparent().remove(e)
# Keep TOC heading text in step with spelling fixes. Page values refreshed after render.
heads=[p.text for p in d.paragraphs if p.style.name.startswith('Heading')]
rows=[p for p in d.paragraphs if '\t' in p.text and p.text[0].isdigit()];assert len(rows)==len(heads)==31
for p,h in zip(rows,heads):p._p.xpath('.//w:t')[0].text=h
out=R/'deliverables/CS395_FinalReport_Efe_Han_Keleş_8October2026.docx';d.save(out)
(R/'sources/changes.json').write_text(json.dumps({'source_sha256':hashlib.sha256(snapshot.read_bytes()).hexdigest(),'protected_sections':['5.1','5.2'],'changes':changes,'pending':['Student text for 5.1 and 5.2','Mentor formal job title is unknown; no title invented','Deadline extension or portal availability not verified'],'supplier_source_verified':'https://www.finnvera.fi/eng/export-and-internationalisation/finnvera-as-an-export-finance-provider/guaranteed-transactions/finnvera-guarantees-nokias-deliveries-to-turk-telekom'},ensure_ascii=False,indent=2)+'\n')
print(out);print('Changed fragments',len(changes),'abstract words',len(d.paragraphs[15].text.split()))
